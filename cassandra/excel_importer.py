

from pathlib import Path
from datetime import datetime

from openpyxl import load_workbook

from cassandra_client import CassandraClient



class ExcelImporter:
    """
    Importateur Excel vers Cassandra.
    """


    def __init__(self):

        self.client = CassandraClient()

        self.base_path = Path("../data")



    # =====================================================
    # Import d'un fichier Excel
    # =====================================================

    def import_file(self, excel_path):

        print("\n" + "=" * 60)
        print(f"Lecture : {excel_path}")
        print("=" * 60)


        workbook = load_workbook(
            excel_path,
            data_only=True
        )


        sheet = workbook.active


        headers = [
            cell.value
            for cell in sheet[1]
            if cell.value is not None
        ]


        imported = 0
        errors = 0



        for row in sheet.iter_rows(
            min_row=2,
            values_only=True
        ):


            if all(value is None for value in row):
                continue



            row = row[:len(headers)]



            data = dict(
                zip(
                    headers,
                    row
                )
            )



            # Conversion timestamp

            if isinstance(
                data.get("timestamp"),
                str
            ):

                try:

                    data["timestamp"] = datetime.strptime(
                        data["timestamp"],
                        "%Y-%m-%d %H:%M:%S"
                    )

                except Exception:

                    pass



            equipment_type = data.get(
                "type"
            )



            try:


                if equipment_type == "ELEVATOR":

                    self.client.insert_elevator(
                        data
                    )


                elif equipment_type == "RMG":

                    self.client.insert_rmg(
                        data
                    )


                elif equipment_type == "RTG":

                    self.client.insert_rtg(
                        data
                    )


                elif equipment_type == "STS":

                    self.client.insert_sts(
                        data
                    )


                elif equipment_type == "STRADDLE":

                    self.client.insert_straddle(
                        data
                    )


                elif equipment_type == "TRACTOR":

                    self.client.insert_tractor(
                        data
                    )


                else:

                    print(
                        f"Type inconnu : {equipment_type}"
                    )

                    continue



                imported += 1



            except Exception as e:


                errors += 1


                print("\n--------------------------------")
                print("Erreur insertion")
                print("--------------------------------")

                print(
                    f"Fichier : {excel_path.name}"
                )

                print(
                    data
                )

                print(
                    e
                )

                print("--------------------------------\n")



        print(
            f"Importées : {imported}"
        )

        print(
            f"Erreurs : {errors}"
        )



    # =====================================================
    # Lancement import complet
    # =====================================================

    def run(self):


        files = [

            self.base_path /
            "elevator" /
            "elevator_data.xlsx",


            self.base_path /
            "rmg" /
            "rmg_data.xlsx",


            self.base_path /
            "rtg" /
            "rtg_data.xlsx",


            self.base_path /
            "sts" /
            "sts_data.xlsx",


            self.base_path /
            "straddle" /
            "straddle_data.xlsx",


            self.base_path /
            "tractor" /
            "tractor_data.xlsx"

        ]



        for file in files:


            if file.exists():

                self.import_file(
                    file
                )

            else:

                print(
                    f"Fichier absent : {file}"
                )



        print(
            "\nImport terminé."
        )


        self.client.close()



# =====================================================

if __name__ == "__main__":


    importer = ExcelImporter()

    importer.run()
